# -*- coding: utf-8 -*-
import re,difflib,unicodedata,docx
from docx.oxml.ns import qn
from html.parser import HTMLParser
def norm(s):
    s=unicodedata.normalize('NFKC',s).lower()
    s=s.replace(chr(8220),'"').replace(chr(8221),'"').replace(chr(8216),"'").replace(chr(8217),"'").replace(chr(8212),'-').replace(chr(8211),'-')
    s=re.sub(r'\s+',' ',s); s=re.sub(r'[^\w\s]','',s,flags=re.UNICODE)
    return s.strip()
class TX(HTMLParser):
    def __init__(s):
        super().__init__(); s.p=[]; s.sk=0
    def handle_starttag(s,t,a):
        if t in('script','style','head'):s.sk+=1
    def handle_endtag(s,t):
        if t in('script','style','head')and s.sk:s.sk-=1
    def handle_data(s,d):
        if not s.sk and d.strip():s.p.append(d.strip()+' ')
def di_sentences(m):
    d=docx.Document(f".tmp-di/docx/mod{m}_di.docx")
    blocks=[]
    pm={p._p:p for p in d.paragraphs}; tm={t._tbl:t for t in d.tables}
    for ch in d.element.body.iterchildren():
        if ch.tag==qn('w:p'):
            p=pm.get(ch)
            if p is not None and p.text.strip(): blocks.append(p.text.strip())
        elif ch.tag==qn('w:tbl'):
            t=tm.get(ch)
            if t: 
                for r in t.rows:
                    for c in r.cells: 
                        if c.text.strip(): blocks.append(c.text.strip())
    sents=[]
    for b in blocks:
        for p in re.split(r'(?<=[.!?:])\s+',b):
            if len(norm(p))>=12: sents.append(p)
    return sents
for m in [1,2,3,4,5]:
    htmltext=open(f"modulo-{m}-conteudo/index.html",encoding='utf-8').read()
    pr=TX(); pr.feed(htmltext); blob=norm(''.join(pr.p))
    sents=di_sentences(m)
    miss=[]
    for s in sents:
        ns=norm(s)
        if ns in blob: continue
        words=ns.split(); anc=max(words,key=len) if words else ''
        best=0; L=len(ns)
        idxs=[mt.start() for mt in re.finditer(re.escape(anc),blob)] if len(anc)>3 else []
        for i in idxs[:40]:
            r=difflib.SequenceMatcher(None,ns,blob[max(0,i-L):i+L]).ratio()
            best=max(best,r)
        if best<0.75: miss.append((round(best,2),s))
    with open(f".tmp-di/mod{m}_verify.txt","w",encoding='utf-8') as f:
        f.write(f"MOD{m}: {len(sents)} sentencas DI | {len(miss)} ainda ausentes (<0.75)\n\n")
        f.write("\n".join(f"[{b}] {s}" for b,s in miss))
    print(f"mod{m}: {len(sents)} DI sentences, {len(miss)} ausentes")
