# -*- coding: utf-8 -*-
import re,unicodedata,docx
from docx.oxml.ns import qn
from html.parser import HTMLParser
def norm(s):
    s=unicodedata.normalize('NFKC',s).lower()
    for a,b in [(chr(8220),'"'),(chr(8221),'"'),(chr(8216),"'"),(chr(8217),"'"),(chr(8212),'-'),(chr(8211),'-')]:
        s=s.replace(a,b)
    s=re.sub(r'\s+',' ',s); s=re.sub(r'[^\w]','',s,flags=re.UNICODE)
    return s.strip()
class TX(HTMLParser):
    def __init__(s):
        super().__init__(); s.p=[]; s.sk=0
    def handle_starttag(s,t,a):
        if t in('script','style','head'): s.sk+=1
    def handle_endtag(s,t):
        if t in('script','style','head') and s.sk: s.sk-=1
    def handle_data(s,d):
        if not s.sk and d.strip(): s.p.append(d.strip()+' ')
d=docx.Document(".tmp-di/docx/mod1_di.docx")
pm={p._p:p for p in d.paragraphs}; tm={t._tbl:t for t in d.tables}
blocks=[]
for ch in d.element.body.iterchildren():
    if ch.tag==qn('w:p'):
        p=pm.get(ch)
        if p is not None and p.text.strip(): blocks.append(p.text.strip())
    elif ch.tag==qn('w:tbl'):
        t=tm.get(ch)
        if t:
            for r in t.rows:
                for c in r.cells:
                    if c.text.strip(): blocks.append(c.text.strip())
sents=[]
for b in blocks:
    for p in re.split(r'(?<=[.!?:])\s+',b):
        if len(norm(p))>=16: sents.append(p)
htmltext=open("modulo-1/index.html",encoding='utf-8').read()
pr=TX(); pr.feed(htmltext); blob=norm(''.join(pr.p))
absent=[]
for s in sents:
    if norm(s) not in blob: absent.append(s)
print(f"GENUINAMENTE AUSENTES (norm não é substring): {len(absent)}")
for s in absent: print("  -",s[:140])
