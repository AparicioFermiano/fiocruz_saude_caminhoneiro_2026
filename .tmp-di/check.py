# -*- coding: utf-8 -*-
import docx, re, difflib, unicodedata
from docx.oxml.ns import qn
def norm(s):
    s=unicodedata.normalize('NFKC',s).lower()
    s=re.sub(r'[^\w]','',s,flags=re.UNICODE)
    return s
# 1) verify PDF vs docx content equality
for m in [1,2,3,4,5]:
    d=docx.Document(f".tmp-di/docx/mod{m}_di.docx")
    dtxt=norm(" ".join(p.text for p in d.paragraphs)+" "+" ".join(c.text for t in d.tables for r in t.rows for c in r.cells))
    ptxt=norm(open(f".tmp-di/mod{m}_pdf.txt",encoding='utf-8').read())
    r=difflib.SequenceMatcher(None,dtxt,ptxt).quick_ratio()
    print(f"mod{m}: docx_chars={len(dtxt)} pdf_chars={len(ptxt)} sim={r:.3f}")
