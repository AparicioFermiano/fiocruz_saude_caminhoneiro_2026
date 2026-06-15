# -*- coding: utf-8 -*-
"""Gera modulo-N-conteudo/index.html fiel ao DI (material/modulo-N-conteudo.pdf == docx)."""
import docx, re, os, html as H
from docx.oxml.ns import qn

PAL = {
 1:("#fecf08","#d4aa00","#524000","#fffde8","#ffe566"),
 2:("#3F56A6","#293872","#FFFFFF","#EEF0FA","#7D96CE"),
 3:("#187848","#0E6238","#FFFFFF","#EAF6EF","#43B64A"),
 4:("#E11B12","#9A1009","#FFFFFF","#FEE9E8","#F47069"),
 5:("#011f4b","#011023","#FFFFFF","#e8edf5","#3a6bc4"),
}
TITLES = {
 1:"O Cuidado da Caminhoneira e do Caminhoneiro na Atenção Primária à Saúde",
 2:"Infecções Sexualmente Transmissíveis e Saúde da População Caminhoneira",
 3:"Saúde Mental da Caminhoneira e do Caminhoneiro",
 4:"Doenças Crônicas Não Transmissíveis na População Caminhoneira",
 5:"Violências, Causas Externas e Promoção do Cuidado nas Estradas",
}
BANNERS = {"SUMÁRIO","APRESENTAÇÃO DO MÓDULO","OBJETIVOS DE APRENDIZAGEM",
 "CARGA HORÁRIA DE ESTUDO","TEXTO-BASE","ENCERRAMENTO DO MÓDULO","REFERÊNCIAS",
 "AUTORAS","AUTORES"}
CALLOUT_WORDS = ("Reflexão","Saiba mais","Saiba Mais","Dica","Atenção","Importante",
 "Lembre-se","Você sabia","Fique atento","Glossário","Para saber mais","Atividade")

# correções de erros óbvios do DI (aprovadas pelo usuário) — aplicadas ao texto exibido
CORRECTIONS = {
 1:[("responsabilidade sanitária esse contexto, compreender","responsabilidade sanitária. Nesse contexto, compreender"),
    ("o vínculo com a equipes de referência","o vínculo com as equipes de referência"),
    ("caminhoneiras e caminheiros","caminhoneiras e caminhoneiros")],
 2:[("diagnosticadas com Infeções Sexualmente","diagnosticadas com Infecções Sexualmente"),
    ("elaboração das autoras ,2026","elaboração das autoras, 2026"),
    ("No caso da caminhoneiras e dos caminhoneiros","No caso das caminhoneiras e dos caminhoneiros")],
 3:[],
 4:[("Entre pessoas camonhoneiras","Entre pessoas caminhoneiras")],
 5:[("conheça mais afundo as orientações","conheça mais a fundo as orientações")],
}
_CORR=[]
def clean(s):
    s=re.sub(r'\s+',' ',s or '').strip()
    for a,b in _CORR: s=s.replace(a,b)
    return s

URL_RE = re.compile(r'(https?://[^\s)\]]+|www\.[^\s)\]]+)')
def linkify(text):
    esc = H.escape(text)
    def repl(m):
        u=m.group(0); u=u.rstrip('.,;')
        tail=m.group(0)[len(u):]
        href=u if u.startswith('http') else 'https://'+u
        return f'<a href="{href}" target="_blank" rel="noopener">{u}</a>{tail}'
    # operate on unescaped then escape carefully: simpler -> escape first, urls have no special chars except &
    return URL_RE.sub(lambda m: repl(m), esc)

def numfmt_map(doc):
    """numId -> 'ol'|'ul' (by level0 numFmt)."""
    res={}
    try:
        npart=doc.part.numbering_part.element
    except Exception:
        return res
    # map abstractNumId -> fmt(level0)
    absfmt={}
    for an in npart.findall(qn('w:abstractNum')):
        aid=an.get(qn('w:abstractNumId'))
        lvl=an.find(qn('w:lvl'))
        fmt=None
        if lvl is not None:
            nf=lvl.find(qn('w:numFmt'))
            if nf is not None: fmt=nf.get(qn('w:val'))
        absfmt[aid]=fmt
    for n in npart.findall(qn('w:num')):
        nid=n.get(qn('w:numId'))
        a=n.find(qn('w:abstractNumId'))
        aid=a.get(qn('w:val')) if a is not None else None
        fmt=absfmt.get(aid)
        res[nid]='ol' if fmt in ('decimal','lowerLetter','upperLetter','lowerRoman','upperRoman') else 'ul'
    return res

def list_info(p):
    ppr=p._p.find(qn('w:pPr'))
    if ppr is None: return None
    npr=ppr.find(qn('w:numPr'))
    if npr is None: return None
    nid=npr.find(qn('w:numId'))
    return nid.get(qn('w:val')) if nid is not None else '?'

def para_imgs(p, doc, outdir, counter):
    out=[]
    for blip in p._p.findall('.//'+qn('a:blip')):
        rid=blip.get(qn('r:embed'))
        if not rid: continue
        try:
            part=doc.part.related_parts[rid]
        except Exception:
            continue
        ext=os.path.splitext(part.partname)[1] or '.png'
        counter[0]+=1
        fn=f"fig{counter[0]:02d}{ext}"
        with open(os.path.join(outdir,'images',fn),'wb') as f:
            f.write(part.blob)
        out.append(fn)
    return out

def render(m):
    global _CORR
    _CORR=CORRECTIONS.get(m,[])
    doc=docx.Document(f".tmp-di/docx/mod{m}_di.docx")
    outdir=f"modulo-{m}-conteudo"
    os.makedirs(os.path.join(outdir,'images'),exist_ok=True)
    nfmt=numfmt_map(doc)
    para_map={p._p:p for p in doc.paragraphs}
    tbl_map={t._tbl:t for t in doc.tables}
    counter=[0]
    # parse TOC section titles for heading detection (collect after detecting SUMÁRIO)
    body=[]
    # gather list of known section heading texts (normalized) from any "n.n" paragraph
    sec_titles=set()
    for p in doc.paragraphs:
        t=clean(p.text)
        mm=re.match(r'^\d+\.\d+\s+(.*)',t)
        if mm:
            base=re.sub(r'\s+\d+\s*$','',mm.group(1)).strip()
            if base: sec_titles.add(base.lower()[:40])
    html_parts=[]
    # iterate body children in order
    def render_para_html(p, in_cell=False):
        t=clean(p.text)
        imgs=para_imgs(p,doc,outdir,counter)
        chunks=[]
        for fn in imgs:
            chunks.append(f'<figure class="fig"><img src="images/{fn}" alt=""></figure>')
        if not t:
            return chunks, None
        low=t.lower()
        # caption
        if re.match(r'^(Quadro|Figura|Tabela)\s+\d+', t) or t.startswith('Fonte:'):
            chunks.append(f'<p class="caption">{linkify(t)}</p>')
            return chunks,'cap'
        # placeholder de imagem -> insere imagem real (recuperada do git em images/)
        mimg=re.search(r'inserir o arquivo (\w+\.\w+)', t)
        if mimg:
            chunks.append(f'<figure class="fig"><img src="images/{mimg.group(1)}" alt=""></figure>')
            return chunks,'img'
        # numbered sub-subsection heading (x.y.z)
        if re.match(r'^\d+\.\d+\.\d+\s', t):
            chunks.append(f'<h3 class="subh">{linkify(t)}</h3>')
            return chunks,'h3'
        # numbered section heading (x.y)
        if re.match(r'^\d+\.\d+\s', t):
            chunks.append(f'<h2 class="sec">{linkify(t)}</h2>')
            return chunks,'h2'
        # unnumbered section heading matched against TOC
        base=re.sub(r'\s+\d+\s*$','',t).strip().lower()[:40]
        if not in_cell and base in sec_titles and len(t)<90:
            chunks.append(f'<h2 class="sec">{linkify(t)}</h2>')
            return chunks,'h2'
        # caso clinico subheading
        if re.match(r'^Caso cl[ií]nico', t):
            chunks.append(f'<h3 class="subh">{linkify(t)}</h3>')
            return chunks,'h3'
        chunks.append(f'<p>{linkify(t)}</p>')
        return chunks,'p'

    pending=[]  # buffered list items (text,type)
    def flush_list():
        nonlocal pending
        if not pending: return
        tag=pending[0][1]
        html_parts.append(f'<{tag}>')
        for txt,_ in pending:
            html_parts.append(f'<li>{linkify(txt)}</li>')
        html_parts.append(f'</{tag}>')
        pending=[]

    def render_table(t):
        rows=t.rows; ncol=len(t.columns)
        # 1x1?
        if len(rows)==1 and len(rows[0].cells)>=1 and all(c.text==rows[0].cells[0].text for c in rows[0].cells):
            cell=rows[0].cells[0]
            celltext=clean(cell.text)
            up=celltext.upper()
            if up in BANNERS or (re.match(r'^[A-ZÀ-Ú0-9ÇÃÕÉÊÍÓÚÂÔ\- ]+$',celltext) and len(celltext)<=42):
                html_parts.append(f'<h1 class="banner">{H.escape(celltext)}</h1>')
                return
            # callout: first paragraph is title
            paras=[pp for pp in cell.paragraphs]
            title=clean(paras[0].text) if paras else celltext
            # title is first CALLOUT word group
            label=None
            for w in CALLOUT_WORDS:
                if title.lower().startswith(w.lower()):
                    label=title[:len(w)]; rest=title[len(w):].strip(); break
            html_parts.append('<aside class="callout">')
            if label:
                html_parts.append(f'<p class="callout-title">{H.escape(label)}</p>')
                start=1
                if rest:
                    paras=list(paras)
                    # render rest as first paragraph
                    html_parts.append(f'<p>{linkify(rest)}</p>')
            else:
                start=0
            # render remaining paragraphs/lists of the cell
            cell_pending=[]
            def cflush():
                nonlocal cell_pending
                if not cell_pending: return
                tg=cell_pending[0][1]
                html_parts.append(f'<{tg}>')
                for tx,_ in cell_pending: html_parts.append(f'<li>{linkify(tx)}</li>')
                html_parts.append(f'</{tg}>')
                cell_pending=[]
            for pp in paras[start:]:
                tx=clean(pp.text)
                if not tx: continue
                li=list_info(pp)
                if li is not None:
                    cell_pending.append((tx, nfmt.get(li,'ul')))
                else:
                    cflush()
                    html_parts.append(f'<p>{linkify(tx)}</p>')
            cflush()
            html_parts.append('</aside>')
            return
        # data table
        html_parts.append('<div class="tbl-wrap"><table>')
        for ri,row in enumerate(rows):
            html_parts.append('<tr>')
            celltag='th' if ri==0 else 'td'
            for c in row.cells:
                txt=clean(c.text)
                html_parts.append(f'<{celltag}>{linkify(txt)}</{celltag}>')
            html_parts.append('</tr>')
        html_parts.append('</table></div>')

    for ch in doc.element.body.iterchildren():
        if ch.tag==qn('w:p'):
            p=para_map.get(ch)
            if p is None: continue
            li=list_info(p)
            t=clean(p.text)
            imgs_present=bool(p._p.findall('.//'+qn('a:blip')))
            if li is not None and t:
                pending.append((t, nfmt.get(li,'ul')))
                continue
            flush_list()
            ch_html,kind=render_para_html(p)
            html_parts.extend(ch_html)
        elif ch.tag==qn('w:tbl'):
            flush_list()
            t=tbl_map.get(ch)
            if t is not None: render_table(t)
    flush_list()

    primary,deep,onp,surface,accent=PAL[m]
    body_html="\n".join(html_parts)
    # resolve placeholder de link do módulo 2 (PDF "Taxas de prevalências IST nacionais e internacionais")
    if m==2:
        body_html=body_html.replace(
            "Disponível em: [inserir link]",
            'Disponível em: <a href="media/taxas_prevalencias_ist.pdf" target="_blank" rel="noopener">Taxas de prevalências IST nacionais e internacionais (PDF)</a>')
    doc_title=f"Módulo {m} — {TITLES[m]}"
    page=f"""<!DOCTYPE html>
<html lang="pt-BR">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>{H.escape(doc_title)}</title>
<link rel="preconnect" href="https://fonts.googleapis.com">
<link href="https://fonts.googleapis.com/css2?family=Baloo+2:wght@500;600;700;800&family=Mukta:wght@300;400;500;600;700&display=swap" rel="stylesheet">
<style>
:root{{--primary:{primary};--deep:{deep};--on:{onp};--surface:{surface};--accent:{accent};}}
*{{box-sizing:border-box}}
body{{margin:0;font-family:'Mukta',system-ui,sans-serif;color:#1f2430;line-height:1.7;background:#fafafa}}
.wrap{{max-width:820px;margin:0 auto;padding:0 1.25rem 5rem}}
header.hero{{background:var(--deep);color:#fff;padding:2.5rem 1.25rem;text-align:center}}
header.hero .kicker{{font-family:'Baloo 2';font-weight:700;letter-spacing:.08em;text-transform:uppercase;font-size:.95rem;opacity:.9}}
header.hero h1{{font-family:'Baloo 2';font-weight:800;margin:.5rem 0 0;font-size:1.9rem;line-height:1.25}}
h1.banner{{font-family:'Baloo 2';font-weight:800;text-transform:uppercase;letter-spacing:.04em;color:var(--deep);font-size:1.4rem;margin:3rem 0 1rem;padding-bottom:.5rem;border-bottom:3px solid var(--primary)}}
h2.sec{{font-family:'Baloo 2';font-weight:700;color:var(--deep);font-size:1.3rem;margin:2.4rem 0 .8rem}}
h3.subh{{font-family:'Baloo 2';font-weight:700;color:var(--primary);font-size:1.12rem;margin:1.8rem 0 .6rem}}
p{{margin:.7rem 0}}
a{{color:var(--deep);word-break:break-word}}
ul,ol{{margin:.7rem 0 .7rem 1.4rem}}
li{{margin:.35rem 0}}
.caption{{font-size:.9rem;color:#555;font-style:italic;margin:.4rem 0 1.2rem}}
.fig{{margin:1.4rem 0;text-align:center}}
.fig img{{max-width:100%;height:auto;border-radius:10px}}
.callout{{background:var(--surface);border-left:5px solid var(--primary);border-radius:0 10px 10px 0;padding:1rem 1.2rem;margin:1.5rem 0}}
.callout-title{{font-family:'Baloo 2';font-weight:700;color:var(--deep);text-transform:uppercase;letter-spacing:.05em;font-size:.95rem;margin:.1rem 0 .5rem}}
.tbl-wrap{{overflow-x:auto;margin:1.4rem 0}}
table{{border-collapse:collapse;width:100%;font-size:.92rem}}
th,td{{border:1px solid #d6dae2;padding:.6rem .7rem;vertical-align:top;text-align:left}}
th{{background:var(--deep);color:#fff;font-family:'Baloo 2'}}
tr:nth-child(even) td{{background:#f4f5f8}}
.authors{{color:#444}}
@media(max-width:600px){{header.hero h1{{font-size:1.45rem}}}}
</style>
</head>
<body>
<header class="hero">
<div class="kicker">Atenção à Saúde da Caminhoneira e do Caminhoneiro na Atenção Primária · Módulo {ROMAN[m]}</div>
<h1>{H.escape(TITLES[m])}</h1>
</header>
<main class="wrap">
{body_html}
</main>
</body>
</html>"""
    with open(os.path.join(outdir,'index.html'),'w',encoding='utf-8') as f:
        f.write(page)
    print(f"modulo-{m}-conteudo/index.html escrito | {len(html_parts)} blocos | {counter[0]} imagens")

ROMAN={1:"I",2:"II",3:"III",4:"IV",5:"V"}
for m in [1,2,3,4,5]:
    render(m)
