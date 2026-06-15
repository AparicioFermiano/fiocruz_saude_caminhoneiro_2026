# -*- coding: utf-8 -*-
import docx
from docx.oxml.ns import qn
from collections import Counter
d=docx.Document(".tmp-di/docx/mod1_di.docx")
# styles
st=Counter(); ex={}
for p in d.paragraphs:
    st[p.style.name]+=1
    if p.text.strip() and p.style.name not in ex: ex[p.style.name]=p.text.strip()[:60]
print("STYLES:")
for s,c in st.most_common(): print(f"  {c:4d} {s!r:28} | {ex.get(s,'')}")
# count lists & numbering ids
fmt=Counter()
for p in d.paragraphs:
    ppr=p._p.find(qn('w:pPr'))
    if ppr is not None:
        npr=ppr.find(qn('w:numPr'))
        if npr is not None:
            ilvl=npr.find(qn('w:ilvl')); numid=npr.find(qn('w:numId'))
            fmt[('list', numid.get(qn('w:val')) if numid is not None else '?')]+=1
print("LIST numIds:", dict(fmt))
# images: count drawings and their position (which paragraph index)
imgs=[]
for i,p in enumerate(d.paragraphs):
    blips=p._p.findall('.//'+qn('a:blip'))
    if blips: imgs.append((i,p.text.strip()[:40]))
print(f"paragraphs with images: {len(imgs)} -> {imgs}")
# hyperlinks count
hl=d.part.rels
print("hyperlink rels:", sum(1 for v in hl.values() if 'hyperlink' in v.reltype))
