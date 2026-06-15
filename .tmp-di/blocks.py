# -*- coding: utf-8 -*-
import docx, re
from docx.oxml.ns import qn
def clean(s): return re.sub(r'\s+',' ',s).strip()
d=docx.Document(".tmp-di/docx/mod1_di.docx")
para_map={p._p:p for p in d.paragraphs}
tbl_map={t._tbl:t for t in d.tables}
out=[]
for ch in d.element.body.iterchildren():
    if ch.tag==qn('w:p'):
        p=para_map.get(ch)
        if p is None: continue
        has_img=bool(p._p.findall('.//'+qn('a:blip')))
        t=clean(p.text)
        # list?
        ppr=p._p.find(qn('w:pPr')); islist=False
        if ppr is not None and ppr.find(qn('w:numPr')) is not None: islist=True
        if t or has_img:
            tag='L' if islist else 'P'
            if has_img: tag='IMG'
            out.append(f"[{tag}] {t}")
    elif ch.tag==qn('w:tbl'):
        out.append("[TBL] (tabela)")
print(f"total blocks: {len(out)}\n")
for i,b in enumerate(out): print(f"{i:3} {b[:150]}")
